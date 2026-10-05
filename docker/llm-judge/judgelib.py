"""Pure helpers for the llm-judge sidecar. Stdlib only (python-docx is imported lazily)."""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path
from typing import Any

TEXT_EXT = {
    ".py", ".cpp", ".h", ".hpp", ".c", ".java", ".js", ".ts", ".md", ".txt", ".json",
    ".yml", ".yaml", ".csv", ".sql", ".html", ".css",
}  # fmt: skip
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
MIN_PAGE_CHARS = 200
DISALLOWED_TOOLS = "Bash,Edit,Write,WebFetch,WebSearch,NotebookEdit,Task,Agent"


def safe_name(name: str, idx: int) -> str:
    """Display/disk name: index prefix + sanitised basename. Never the raw filename."""
    base = re.split(r"[\\/]", name)[-1]
    return f"{idx:02d}_" + re.sub(r"[^\w.\-]+", "_", base)[:100]


def pdf_needs_visual(page_texts: list[str]) -> bool:
    return any(len(re.sub(r"\s", "", p)) < MIN_PAGE_CHARS for p in page_texts)


def _pdf_text(path: Path) -> tuple[str | None, bool]:
    proc = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, timeout=60
    )
    if proc.returncode != 0:
        return None, True
    pages = proc.stdout.split("\f")
    if pages and not pages[-1].strip():
        pages.pop()
    text = proc.stdout.replace("\f", "\n").strip()
    return (text or None), pdf_needs_visual(pages) if pages else True


def _docx_text(path: Path) -> str | None:
    import docx  # python3-docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts).strip() or None


def extract(path: Path, mime: str) -> tuple[str | None, bool]:
    """(text or None, offer_for_read)."""
    ext = path.suffix.lower()
    if ext == ".pdf" or mime == "application/pdf":
        return _pdf_text(path)
    if ext == ".docx":
        try:
            return _docx_text(path), False
        except Exception:  # corrupt/zip-bomb-ish file: treat as unreadable
            return None, False
    if ext in TEXT_EXT:
        return path.read_text(encoding="utf-8", errors="replace"), False
    if ext in IMAGE_EXT:
        return None, True
    return None, False


def build_prompt_with_files(
    prompt: str, files: list[tuple[str, str | None, bool]], root: Path = Path("/tmp/judge")
) -> str:
    """Append every file after the unchanged prompt (stable prefix keeps prompt caching).

    ``files`` items are (safe_name, text, offer_for_read); the file lives at
    ``root / safe_name`` (root = the request's work dir) for the Read-tool pointer.
    """
    out = [prompt]
    for name, text, offered in files:
        out.append(f"\n--- FILE: {name} ---\n")
        if text is not None:
            out.append(text)
        elif offered:
            out.append("[no text layer]")
        else:
            out.append("[unsupported file type]")
        if offered:
            out.append(f"\n[You may open this file with the Read tool: {root / name}]")
    return "".join(out)


def claude_argv(model: str, system: str, workdir: Path, allow_read: bool) -> list[str]:
    argv = [
        "claude", "-p", "--output-format", "json", "--model", model,
        "--append-system-prompt", system, "--max-turns", "6", "--add-dir", str(workdir),
    ]  # fmt: skip
    if allow_read:
        argv += ["--allowedTools", "Read"]
    return argv + ["--disallowedTools", DISALLOWED_TOOLS]


def parse_cli_output(stdout: str) -> dict[str, Any]:
    try:
        data = json.loads(stdout)
    except ValueError as exc:
        raise ValueError("claude output is not JSON") from exc
    if isinstance(data, list):  # some versions emit an event array; take the result event
        data = next((e for e in reversed(data) if isinstance(e, dict) and "result" in e), {})
    if not isinstance(data, dict) or not isinstance(data.get("result"), str):
        raise ValueError("claude output has no string 'result'")
    if data.get("is_error"):
        raise ValueError(f"claude reported an error: {data['result'][:300]}")
    usage = data.get("usage")
    used = data.get("modelUsage")
    model = data.get("model") or (next(iter(used)) if isinstance(used, dict) and used else "")
    return {
        "result": data["result"],
        "model": model,
        "usage": usage if isinstance(usage, dict) else {},
    }
